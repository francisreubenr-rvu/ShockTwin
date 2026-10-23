#!/usr/bin/env python3
"""Build ShockTwin's native, editable 14-slide pitch and deterministic previews.

Requires the genuine background images and app captures. Fails closed when an
input is missing; no visual placeholder or fabricated app screen is generated.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
import shutil
import subprocess
import sys

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_TICK_LABEL_POSITION, XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentations'
FONTS = OUT / 'fonts'
ASSETS = OUT / 'composition'
W, H = 20.0, 11.25
PX = 96
AMBER = 'FFB05C'
WHITE = 'FFFFFF'
SOFT = 'E5EBE9'
DARK = '0B2329'
FONT = {'serif': 'Instrument Serif', 'body': 'Inter', 'mono': 'DM Mono', 'pixel': 'Silkscreen'}
FONT_FILES = {
    'serif': FONTS / 'InstrumentSerif-Regular.ttf',
    'italic': FONTS / 'InstrumentSerif-Italic.ttf',
    'body': FONTS / 'Inter-Regular.ttf',
    'mono': FONTS / 'DMMono-Regular.ttf',
    'pixel': FONTS / 'Silkscreen-Regular.ttf',
}
DATA = json.loads((ROOT / 'docs' / 'presentation.json').read_text())
LAYOUT_LOG: list[dict] = []
BACKGROUND_HASHES: list[str] = []


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font_at(role: str, pt: float):
    return ImageFont.truetype(str(FONT_FILES[role]), round(pt * PX / 72))


def length(text: str, role: str, pt: float) -> float:
    return font_at(role, pt).getlength(text) / PX


def wrap(text: str, role: str, pt: float, width: float) -> list[str]:
    output = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = (line + ' ' + word).strip()
            if line and length(candidate, role, pt) > width:
                output.append(line)
                line = word
            else:
                line = candidate
        output.append(line)
    return output


def alpha(color, opacity: float):
    node = color._color._xClr
    for child in list(node):
        if child.tag.endswith('}alpha'):
            node.remove(child)
    opacity_node = OxmlElement('a:alpha')
    opacity_node.set('val', str(round(opacity * 100000)))
    node.append(opacity_node)


def set_fill(shape, color: str, opacity: float = 1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    alpha(shape.fill.fore_color, opacity)


def set_line(shape, color=WHITE, width=1.5, opacity=.45):
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    alpha(shape.line.color, opacity)


def box(slide, x, y, w, h, text='', *, font='body', size=22, color=WHITE,
        bold=False, italic=False, align=PP_ALIGN.LEFT, margin=0,
        leading=1.12, role='text', wrap_text=True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = f'{role}: {text[:70]}'
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.word_wrap = wrap_text
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = wrap(text, font, size, w - margin * 2) if wrap_text else text.split('\n')
    for i, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(size * leading)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT[font]
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.bold = bold
        run.font.italic = italic
    estimate = len(lines) * size * leading / 72
    if estimate > h - margin * 2 + .08:
        raise ValueError(f'Text height overflow slide {len(prs.slides)} {role}: {estimate:.2f}>{h:.2f}: {text}')
    for line in lines:
        if length(line, font, size) > w - 2 * margin + .08:
            raise ValueError(f'Text width overflow slide {len(prs.slides)} {role}: {line}')
    LAYOUT_LOG.append({'slide': len(prs.slides), 'role': role, 'text': text,
                       'box': [x, y, w, h], 'font': FONT[font], 'size': size,
                       'lines': lines, 'estimatedHeight': round(estimate, 3)})
    return shape


def connector(slide, x1, y1, x2, y2, color=WHITE, width=1.5, opacity=.8):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    set_line(line, color, width, opacity)
    return line


def callout(slide, dot, bend, end):
    connector(slide, *dot, *bend)
    connector(slide, *bend, *end)
    x, y = dot
    outer = slide.shapes.add_shape(SHAPE.OVAL, Inches(x-.105), Inches(y-.105), Inches(.21), Inches(.21))
    outer.fill.background()
    set_line(outer, WHITE, 1.2, .85)
    inner = slide.shapes.add_shape(SHAPE.OVAL, Inches(x-.037), Inches(y-.037), Inches(.074), Inches(.074))
    set_fill(inner, WHITE)
    inner.line.fill.background()


def prepare_background(number: int) -> Image.Image:
    source = OUT / 'backgrounds' / f'slide-{number:02}.png'
    BACKGROUND_HASHES.append(sha256(source))
    cached=ASSETS / f'background-{number:02}.png'
    if cached.exists() and cached.stat().st_mtime >= max(source.stat().st_mtime, Path(__file__).stat().st_mtime):
        return Image.open(cached).convert('RGB')
    background = ImageOps.fit(Image.open(source).convert('RGB'), (1920, 1080), method=Image.Resampling.LANCZOS)
    # Deep left scrim for hall readability, preserving the painterly subject at right.
    layer = Image.new('RGBA', background.size)
    pixels = []
    for yy in range(1080):
        vertical = max(0., 1-yy/175) * .25 + max(0., (yy-850)/230) * .42
        for xx in range(1920):
            horizontal = (.56/(1+math.exp((xx-1000)/160))) if number in (6,11,12,13,14) else .44 * max(0., 1 - xx/1190) ** .75
            vignette = .08 * ((xx-960)/960) ** 2
            right_scrim = .32 * max(0., min(1., (xx-850)/450)) if number in (11,13) else 0.
            opacity = min(.75, horizontal + vertical + vignette + right_scrim)
            pixels.append((4, 17, 23, round(255*opacity)))
    layer.putdata(pixels)
    composed = Image.alpha_composite(background.convert('RGBA'), layer).convert('RGB')
    composed.save(ASSETS / f'background-{number:02}.png')
    return composed


def glass(slide, background: Image.Image, x, y, w, h, *, tag='glass', border=WHITE,
          opacity=.16, dark=False, radius=.25):
    crop = background.crop((round(x*PX), round(y*PX), round((x+w)*PX), round((y+h)*PX))).convert('RGBA')
    crop = crop.filter(ImageFilter.GaussianBlur(18))
    tint = Image.new('RGBA', crop.size, ((5, 19, 27) if dark else (255, 255, 255)) + (round(255*opacity),))
    crop = Image.alpha_composite(crop, tint)
    mask = Image.new('L', crop.size)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, crop.width-1, crop.height-1), radius=round(radius*PX), fill=255)
    crop.putalpha(mask)
    safe_tag=re.sub(r'[^a-zA-Z0-9_-]', '-', tag)
    path = ASSETS / f'{len(prs.slides):02}-{safe_tag}-{len(slide.shapes)}.png'
    crop.save(path)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    border_shape = slide.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    border_shape.name = f'Native editable glass border: {tag}'
    border_shape.adjustments[0] = min(.5, radius/min(w, h))
    border_shape.fill.background()
    set_line(border_shape, border, 1.5, .45 if border == WHITE else .85)
    return border_shape


def chrome(slide, background, item):
    frame = slide.shapes.add_shape(SHAPE.RECTANGLE, Inches(.35), Inches(.35), Inches(19.30), Inches(10.55))
    frame.fill.background()
    set_line(frame, WHITE, .65, .45)
    # Native freeform left semicircle with a native circular outline.
    circle = slide.shapes.add_shape(SHAPE.OVAL, Inches(.72), Inches(.64), Inches(.30), Inches(.30))
    circle.fill.background()
    set_line(circle, WHITE, .9, .95)
    cx, cy, radius = .87, .79, .15
    points = [(round(cx*1000), round(cy*1000))]
    points += [(round((cx+radius*math.cos(a))*1000), round((cy+radius*math.sin(a))*1000)) for a in [math.pi/2+i*math.pi/24 for i in range(25)]]
    builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=Inches(.001))
    builder.add_line_segments(points[1:], close=True)
    half = builder.convert_to_shape()
    set_fill(half, WHITE)
    half.line.fill.background()
    box(slide, 1.18, .65, 3.4, .32, 'SHOCKTWIN', font='body', size=18, bold=True, role='wordmark')
    for name, x, width in [('Problem', 7.0, 1.35), ('ShockTwin', 8.85, 1.75), ('SAP', 11.1, .65), ('Impact', 12.25, 1.3)]:
        box(slide, x, .69, width, .3, name, font='mono', size=13, align=PP_ALIGN.CENTER, role='navigation')
        if item['section'] == name:
            connector(slide, x, 1.11, x+width, 1.11, AMBER, 1.7, 1)
    glass(slide, background, 18.03, .59, 1.22, .46, tag='page-counter', radius=.16, opacity=.12)
    box(slide, 18.10, .69, 1.08, .28, f"{item['number']:02} / 14", font='mono', size=13, align=PP_ALIGN.CENTER, role='page counter', wrap_text=False)
    connector(slide, .7, 1.25, 19.3, 1.25, width=.55, opacity=.3)
    box(slide, .72, 10.58, 4.6, .24, 'SAP HACKFEST 2026 · RV UNIVERSITY', font='mono', size=9.5, role='event footer', wrap_text=False)
    box(slide, 15.25, 10.58, 4.02, .24, 'RESILIENT SUPPLY CHAINS', font='mono', size=9.5, align=PP_ALIGN.RIGHT, role='track footer', wrap_text=False)
    box(slide, 5.1, 10.48, 9.9, .42, item['sourceLine'], font='mono', size=8.7, align=PP_ALIGN.CENTER, leading=1.18, role='source footer')


def eyebrow(slide, background, text):
    width = min(9.1, length(text, 'mono', 13) + .72)
    glass(slide, background, .8, 1.55, width, .46, tag='eyebrow', radius=.22, opacity=.14)
    dot = slide.shapes.add_shape(SHAPE.OVAL, Inches(1.0), Inches(1.73), Inches(.08), Inches(.08))
    set_fill(dot, AMBER)
    dot.line.fill.background()
    box(slide, 1.21, 1.68, width-.55, .26, text.upper(), font='mono', size=13, role='eyebrow', wrap_text=False)


HEADLINES = {
    1: (112, ['See the shock', 'before it hits.']),
    2: (108, ['One fire.', '$210 billion lost.']),
    3: (104, ["Most firms can't", 'see past tier 1.']),
    4: (100, ['The danger sits', 'below the surface.']),
    5: (108, ['The Red Sea', 'closed.', 'India paid.']),
    6: (102, ['Toyota survived', 'because it', 'had a map.']),
    7: (146, ['Meet', 'ShockTwin.']),
    8: (100, ['Find the suppliers', 'you never knew.']),
    9: (100, ['War-game', 'the storm', 'before it lands.']),
    10: (104, ['Borrow capacity', 'in minutes,', 'not weeks.']),
    11: (132, ['Built on', 'SAP BTP.']),
    12: (102, ['Designed step', 'by step', 'with empathy.']),
    13: (112, ['What resilience', 'is worth.']),
    14: (110, ['Never be', 'surprised again.']),
}


def headline(slide, background, item):
    number = item['number']
    size, lines = HEADLINES[number]
    assert ' '.join(lines) == item['headline']
    x, y, w, h = .8, 2.3, 9.1, 4.3
    if number == 2:
        glass(slide, background, .63, 2.17, 9.55, 3.78, tag='overall-shortage-headline', dark=True, opacity=.18)
        box(slide, .88, 2.36, 8.8, .28, 'OVERALL CHIP SHORTAGE · 2021', font='mono', size=13, role='headline context')
        y = 2.85
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = 'Editable Instrument Serif headline'
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    full = item['headline']
    accent_start = full.index(item['accent'])
    accent_end = accent_start + len(item['accent'])
    cursor = 0
    for line_index, line in enumerate(lines):
        p = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = Pt(size*.94)
        line_width = 0.
        for word_index, word in enumerate(line.split(' ')):
            position = full.index(word, cursor)
            is_accent = position < accent_end and position+len(word) > accent_start
            text = word + (' ' if word_index < len(line.split(' '))-1 else '')
            run = p.add_run()
            run.text = text
            run.font.name = FONT['serif']
            run.font.size = Pt(size)
            run.font.italic = is_accent
            run.font.color.rgb = RGBColor.from_string(AMBER if is_accent else WHITE)
            line_width += length(text, 'italic' if is_accent else 'serif', size)
            cursor = position + len(word)
        if line_width > w+.06:
            raise ValueError(f'Headline {number} width overflow: {line_width:.2f}>{w}: {line}')
    LAYOUT_LOG.append({'slide': number, 'role': 'headline', 'text': item['headline'], 'font': FONT['serif'], 'size': size, 'accent': item['accent'], 'box': [x,y,w,h], 'lines': lines})
    box(slide, .85, 6.67, 8.8, 1.11, item['support'], font='body', size=23, leading=1.15, role='single supporting sentence')


def source(slide, x, y, w, text, size=10, h=.40):
    return box(slide, x, y, w, h, text, font='mono', size=size, color=SOFT, leading=1.12, role='card source')


def stat_card(slide, background, x, y, w, h, stat, *, value=None, size=48, pixel=True, label=None, source_text=None, dark=False):
    glass(slide, background, x,y,w,h, tag=stat['label'].replace(' ', '-').lower(), dark=dark, opacity=.20 if dark else .16)
    box(slide,x+.23,y+.18,w-.46,.50,label or stat['label'],font='mono',size=13,leading=1.07,role='statistic label')
    value = (value if value is not None else stat['value']).replace('−','-')
    role = 'pixel' if pixel else 'body'
    lines = wrap(value, role, size, w-.46)
    available = h-1.0
    while (len(lines)*size*1.05/72 > available or max(length(line, role, size) for line in lines)>w-.46) and size>22:
        size-=1
        lines=wrap(value,role,size,w-.46)
    box(slide,x+.23,y+.70,w-.46,available,value,font=role,size=size,leading=1.05,role='editable statistic')
    source(slide,x+.23,y+h-.48,w-.46,source_text or stat['source'],10 if w>4 else 8.5,h=.39)


def device(slide, background, screenshot, x,y,w,h,label, *, crop=None):
    glass(slide,background,x,y,w,h,tag='real-app-device',dark=True,opacity=.18)
    box(slide,x+.24,y+.20,w-.48,.35,label,font='mono',size=13,color=WHITE,role='real screenshot label')
    path = ROOT / screenshot
    image = Image.open(path).convert('RGB')
    if crop:
        image = image.crop(crop)
    inner_x,inner_y,inner_w,inner_h = x+.22,y+.70,w-.44,h-1.28
    scale = min(inner_w/image.width,inner_h/image.height)
    draw_w,draw_h=image.width*scale,image.height*scale
    actual_x=inner_x+(inner_w-draw_w)/2
    actual_y=inner_y+(inner_h-draw_h)/2
    clipped=ASSETS/f"slide-{len(prs.slides):02}-real-{Path(screenshot).stem}.png"
    image.save(clipped)
    slide.shapes.add_picture(str(clipped),Inches(actual_x),Inches(actual_y),Inches(draw_w),Inches(draw_h))
    source(slide,x+.24,y+h-.39,w-.48,'Illustrative local app · actual browser capture',9.2,h=.28)
    LAYOUT_LOG.append({'slide':len(prs.slides),'role':'genuine screenshot','source':str(path.relative_to(ROOT)),'sourceSha256':sha256(path),'crop':crop,'placedBox':[actual_x,actual_y,draw_w,draw_h]})


def clear_chart_background(chart):
    for parent in [chart._chartSpace, chart._chartSpace.chart.plotArea]:
        sppr=parent.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
        if sppr is None:
            sppr=OxmlElement('c:spPr');parent.append(sppr)
        fill=OxmlElement('a:noFill');sppr.append(fill)
        line=OxmlElement('a:ln');line.append(OxmlElement('a:noFill'));sppr.append(line)


def chart_theme(chart, *, max_value, major_unit, number_format='0', category_skip=1):
    chart.chart_style=10
    chart.has_title=False
    chart.has_legend=False
    clear_chart_background(chart)
    for axis in [chart.category_axis,chart.value_axis]:
        axis.tick_labels.font.name=FONT['mono']
        axis.tick_labels.font.size=Pt(12)
        axis.tick_labels.font.color.rgb=RGBColor.from_string(WHITE)
        axis.format.line.color.rgb=RGBColor.from_string(WHITE)
        axis.format.line.width=Pt(.6)
        axis.major_tick_mark=XL_TICK_MARK.NONE
        axis.minor_tick_mark=XL_TICK_MARK.NONE
    chart.value_axis.minimum_scale=0
    chart.value_axis.maximum_scale=max_value
    chart.value_axis.major_unit=major_unit
    chart.value_axis.tick_labels.number_format=number_format
    chart.value_axis.has_major_gridlines=True
    chart.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string(WHITE)
    chart.value_axis.major_gridlines.format.line.width=Pt(.5)
    alpha(chart.value_axis.major_gridlines.format.line.color,.18)
    chart.category_axis.has_major_gridlines=False
    if category_skip>1:
        ax=chart.category_axis._element
        for key in ['tickLblSkip','tickMarkSkip']:
            node=OxmlElement('c:'+key);node.set('val',str(category_skip));ax.append(node)


def visibility_chart(slide,background):
    glass(slide,background,10.45,5.65,8.65,4.18,tag='native-visibility-chart')
    box(slide,10.73,5.9,8.05,.35,'VISIBILITY BEYOND TIER ONE',font='mono',size=14,role='chart title')
    data=CategoryChartData();data.categories=['2022','2023','2024'];data.add_series('Visibility',[56,37,30])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(10.73),Inches(6.48),Inches(7.98),Inches(2.70),data).chart
    chart_theme(chart,max_value=60,major_unit=20,number_format='0"%"')
    series=chart.series[0]
    series.format.line.color.rgb=RGBColor.from_string(WHITE);series.format.line.width=Pt(2.5)
    series.marker.style=XL_MARKER_STYLE.CIRCLE;series.marker.size=10
    series.marker.format.fill.solid();series.marker.format.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    series.marker.format.line.fill.background()
    series.has_data_labels=True
    labels=series.data_labels;labels.show_value=True;labels.position=XL_LABEL_POSITION.ABOVE;labels.number_format='0"%"'
    labels.font.name=FONT['pixel'];labels.font.size=Pt(26);labels.font.color.rgb=RGBColor.from_string(WHITE)
    for index, point in enumerate(series.points):
        point.data_label.font.name=FONT['pixel'];point.data_label.font.size=Pt(26)
        point.data_label.font.color.rgb=RGBColor.from_string(AMBER if index==2 else WHITE)
    point=series.points[2];point.marker.style=XL_MARKER_STYLE.CIRCLE;point.marker.size=11
    point.marker.format.fill.solid();point.marker.format.fill.fore_color.rgb=RGBColor.from_string(AMBER);point.marker.format.line.fill.background()
    decisive=slide.shapes.add_shape(SHAPE.OVAL,Inches(17.325-.075),Inches(7.89-.075),Inches(.15),Inches(.15))
    set_fill(decisive,AMBER);decisive.line.fill.background();decisive.name='Editable amber latest visibility marker'
    source(slide,10.73,9.37,8.0,'McKinsey Supply Chain Risk Survey 2024',10)


def redsea_chart(slide,background):
    glass(slide,background,10.45,5.65,8.65,4.18,tag='native-container-cost-chart')
    box(slide,10.73,5.90,8.0,.35,'INDIA–EUROPE CONTAINER COST',font='mono',size=14,role='chart title')
    data=CategoryChartData();data.categories=['Before','Disrupted'];data.add_series('Cost',[600,1500])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(10.73),Inches(6.48),Inches(7.98),Inches(2.70),data).chart
    chart_theme(chart,max_value=2000,major_unit=1000,number_format='$#,##0')
    plot=chart.plots[0];plot.gap_width=100
    series=chart.series[0]
    for i,color in enumerate([WHITE,AMBER]):
        series.points[i].format.fill.solid();series.points[i].format.fill.fore_color.rgb=RGBColor.from_string(color)
        series.points[i].format.line.fill.background()
    # Native editable direct labels avoid LibreOffice wrapping currency inside its narrow auto-label boxes.
    # The native bar chart workbook still contains the exact source values 600 and 1500.
    box(slide,12.10,7.70,2.40,.46,'$600',font='pixel',size=25,align=PP_ALIGN.CENTER,role='direct native bar value')
    box(slide,15.35,6.73,3.00,.46,'$1,500',font='pixel',size=25,color=AMBER,align=PP_ALIGN.CENTER,role='direct native bar value')
    source(slide,10.73,9.37,8.0,'Reuters, Jan 2024 · approximate costs',10)


def simulation_chart(slide,background,result):
    x,y,w,h=.85,7.98,8.65,2.24
    glass(slide,background,x,y,w,h,tag='native-real-simulation-chart')
    box(slide,x+.2,y+.12,w-.4,.27,'RESIN STOCK LEFT · ILLUSTRATIVE',font='mono',size=12,role='chart title')
    data=CategoryChartData();data.categories=[str(point['day']) for point in result['baselineSeries']]
    data.add_series('Without action',[point['stock'] for point in result['baselineSeries']])
    data.add_series('With day-3 switch',[point['stock'] for point in result['mitigatedSeries']])
    chart=slide.shapes.add_chart(XL_CHART_TYPE.LINE,Inches(x+.16),Inches(y+.48),Inches(w-.32),Inches(1.45),data).chart
    chart_theme(chart,max_value=1800,major_unit=900,category_skip=6)
    chart.category_axis.tick_labels.font.size=Pt(9)
    chart.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.NONE
    chart.value_axis.tick_labels.font.size=Pt(9)
    plot_area=chart._chartSpace.chart.plotArea
    layout=OxmlElement('c:layout');plot_area.insert(0,layout)
    manual=OxmlElement('c:manualLayout')
    for key,value in [('layoutTarget','inner'),('xMode','edge'),('yMode','edge'),('wMode','factor'),('hMode','factor'),('x','.08'),('y','.10'),('w','.90'),('h','.70')]:
        child=OxmlElement('c:'+key);child.set('val',value);manual.append(child)
    layout.append(manual)
    for day in [0,6,12,18,24,30]:
        box(slide,1.676+((day+.5)/31)*7.497-.14,9.71,.28,.22,str(day),font='mono',size=9,align=PP_ALIGN.CENTER,role='native editable sparse day tick')
    box(slide,7.55,8.10,1.60,.32,'DAY 18',font='pixel',size=18,color=AMBER,align=PP_ALIGN.RIGHT,role='native failure readout')
    for i,series in enumerate(chart.series):
        series.format.line.color.rgb=RGBColor.from_string(WHITE);series.format.line.width=Pt(2)
        series.marker.style=XL_MARKER_STYLE.NONE
        if i:series.format.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    point=chart.series[0].points[18]
    point.marker.style=XL_MARKER_STYLE.CIRCLE;point.marker.size=8;point.marker.format.fill.solid();point.marker.format.fill.fore_color.rgb=RGBColor.from_string(AMBER);point.marker.format.line.fill.background()
    box(slide,5.9,8.50,3.0,.25,'DAY-3 SWITCH',font='mono',size=10,role='direct curve label')
    box(slide,5.25,9.25,3.9,.25,'DAY 18 · STOCK-OUT',font='mono',size=10,color=AMBER,role='direct failure label')
    marker=slide.shapes.add_shape(SHAPE.OVAL,Inches(6.150-.05),Inches(9.62-.05),Inches(.10),Inches(.10))
    set_fill(marker,AMBER);marker.line.fill.background();marker.name='Editable amber day-18 failure marker'
    source(slide,x+.2,y+h-.25,w-.4,'Real local API · 100 units/day · 30-day horizon',8.5,h=.22)


def fill_slide(slide,background,item,result):
    n=item['number'];stats=item['stats']
    if n==1:
        glass(slide,background,.85,8.15,8.65,1.82,tag='product')
        box(slide,1.10,8.35,8.1,.3,'PRODUCT',font='mono',size=13,role='card label')
        box(slide,1.10,8.82,8.1,.57,'ShockTwin',font='body',size=36,role='product name')
        source(slide,1.1,9.57,8.0,'Detect · Simulate · Respond · supplied brief',10)
        glass(slide,background,12.05,8.30,7.05,1.67,tag='team')
        box(slide,12.30,8.49,6.55,.3,'TEAM · RV UNIVERSITY',font='mono',size=13,role='card label')
        box(slide,12.30,8.94,6.55,.70,'Francis Reuben R · Avaneesh\nAffan · Vivek',font='body',size=22,leading=1.12,role='team')
        callout(slide,(15.5,4.95),(11.2,7.68),(9.28,8.28))
    elif n==2:
        stat_card(slide,background,.85,8.15,4.22,1.95,stats[1],size=48)
        stat_card(slide,background,5.32,8.15,4.22,1.95,stats[2],size=48)
    elif n==3:
        visibility_chart(slide,background)
    elif n==4:
        stat_card(slide,background,.85,8.15,4.22,1.95,stats[1],size=47)
        stat_card(slide,background,5.32,8.15,4.22,1.95,stats[2],size=39)
        stat_card(slide,background,10.45,7.60,5.10,2.45,stats[0],size=64)
        callout(slide,(17.3,6.85),(16.4,7.25),(15.35,7.9))
    elif n==5:
        stat_card(slide,background,.85,8.15,4.22,1.95,stats[1],size=48)
        stat_card(slide,background,5.32,8.15,4.22,1.95,stats[2],size=46)
        redsea_chart(slide,background)
    elif n==6:
        for i,stat in enumerate(stats):stat_card(slide,background,.85+i*6.17,8.05,5.92,2.05,stat,size=48)
        callout(slide,(15.1,5.0),(10.2,7.36),(6.50,8.18))
    elif n==7:
        for i,stat in enumerate(stats):
            x=.85+i*6.17
            glass(slide,background,x,8.00,5.92,2.10,tag='product-phase')
            box(slide,x+.26,8.25,5.4,.4,stat['label'],font='pixel',size=22,role='product ordinal')
            box(slide,x+.26,8.85,5.4,.70,stat['value'],font='body',size=35,role='product phase')
            source(slide,x+.26,9.65,5.3,'Supplied product loop',10)
    elif n==8:
        device(slide,background,item['screenshotPaths'][0],10.45,2.30,8.65,5.0,'DETECT · REAL APP · ILLUSTRATIVE',crop=(270,353,1588,892))
        labels=['IMMEDIATE\nSUPPLIERS','SHARED\nUPSTREAM PLANT','DEPENDENT\nSUPPLIERS']
        for i,stat in enumerate(stats):stat_card(slide,background,.85+i*2.98,8.15,2.76,1.95,stat,size=43,label=labels[i],source_text='Illustrative seed')
    elif n==9:
        device(slide,background,item['screenshotPaths'][0],10.45,2.30,8.65,6.4,'SIMULATE · REAL APP · ILLUSTRATIVE')
        simulation_chart(slide,background,result)
    elif n==10:
        device(slide,background,item['screenshotPaths'][0],10.45,2.3,8.65,3.65,'RANKED CAPACITY · ILLUSTRATIVE',crop=(270,400,1882,1019))
        device(slide,background,item['screenshotPaths'][1],10.45,6.20,8.65,3.65,'HUMAN APPROVED · MOCK ERP',crop=(270,353,1882,935))
        configs=[('40%','UP TO · EMPTY RUNNING','DPIIT–NCAER, 2025'),('74.6%','CAPACITY UTILISATION','RBI OBICUS Q1 2026–27 [VERIFY]'),('~60,000','VENTILATORS · ~3 MONTHS','COVID sharing case [source VERIFY]')]
        for i,(value,label,src) in enumerate(configs):stat_card(slide,background,.85+i*2.98,8.15,2.76,1.95,stats[i],value=value,label=label,source_text=src,size=36)
    elif n==11:
        for i,stat in enumerate(stats):
            x,y,w,h=10.45,2.35+i*1.06,8.65,.91
            glass(slide,background,x,y,w,h,tag='SAP-'+stat['label'],border=AMBER if stat['label']=='AI' else WHITE,dark=True,opacity=.20)
            box(slide,x+.24,y+.19,1.30,.30,stat['label'],font='mono',size=14,color=AMBER if stat['label']=='AI' else WHITE,role='architecture layer')
            box(slide,x+1.73,y+.13,6.65,.39,stat['value'],font='body',size=22,role='architecture component')
            source(slide,x+1.73,y+.62,6.65,stat['source'],9.0,h=.24)
    elif n==12:
        for i,stat in enumerate(stats):
            x=.85+i*3.69
            glass(slide,background,x,8.0,3.44,2.1,tag=stat['label'],dark=True,opacity=.20)
            box(slide,x+.20,8.19,3.04,.30,stat['label'],font='mono',size=12.7,role='design thinking phase')
            box(slide,x+.20,8.72,3.04,.88,stat['value'].replace('Hidden dependency + response gap','Hidden dependency\n+ response gap'),font='body',size=22,leading=1.05,role='design thinking status')
            source(slide,x+.20,9.77,3.04,'Project process record',8.5,h=.23)
    elif n==13:
        for i,stat in enumerate(stats):
            x=10.45+(i%2)*4.48;y=4.15+(i//2)*2.92
            stat_card(slide,background,x,y,4.17,2.65,stat,size=45 if i!=3 else 32,source_text=stat['source'].replace('Collaboration study supplied in brief','Collaboration study from brief'),dark=True)
    elif n==14:
        stat_card(slide,background,.85,8.15,8.65,1.95,stats[0],size=31,pixel=False)
        glass(slide,background,10.45,8.15,8.65,1.95,tag='team-project-id')
        box(slide,10.70,8.34,5.55,.3,'TEAM',font='mono',size=13,role='team label')
        box(slide,10.70,8.80,5.55,.75,'Francis Reuben R · Avaneesh\nAffan · Vivek',font='body',size=21.5,role='team')
        box(slide,16.50,8.34,2.3,.3,'P-ID',font='mono',size=13,role='project id label')
        box(slide,16.50,8.91,2.3,.4,'P2012988948',font='mono',size=19,role='project id',wrap_text=False)
        source(slide,10.70,9.71,8.0,'Team and P-ID: supplied brief',10,h=.25)


def notes(slide,item,start):
    end=start+item['duration']
    timestamp=lambda seconds:f'{seconds//60}:{seconds%60:02}'
    frame=slide.notes_slide.notes_text_frame
    frame.text=(f"SLIDE {item['number']:02} · {item['duration']} SECONDS · {timestamp(start)}–{timestamp(end)}\n\n"
                +item['notes']+'\n\nSOURCE / QUALIFICATION (not additional spoken copy):\n'+item['sourceLine'])
    return end


def input_check():
    missing=[]
    for number in range(1,15):
        path=OUT/'backgrounds'/f'slide-{number:02}.png'
        if not path.is_file():missing.append(str(path.relative_to(ROOT)))
    for item in DATA:
        for screenshot in item['screenshotPaths']:
            if not (ROOT/screenshot).is_file():missing.append(screenshot)
    required=[ROOT/'docs/screenshots/simulation-result.json',*FONT_FILES.values()]
    for path in required:
        if not path.is_file():missing.append(str(path.relative_to(ROOT)))
    if missing:raise FileNotFoundError('Missing genuine production inputs:\n'+'\n'.join(missing))
    if sum(item['duration'] for item in DATA)!=450 or len(DATA)!=14:raise ValueError('Presentation timing/content contract is invalid')


def validate_file(path):
    import zipfile
    with zipfile.ZipFile(path) as archive:
        names=archive.namelist()
        chart_parts=[name for name in names if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
        note_parts=[name for name in names if name.startswith('ppt/notesSlides/notesSlide') and name.endswith('.xml')]
        if len(chart_parts)!=3 or len(note_parts)!=14:raise ValueError(f'Native chart/notes count invalid: {len(chart_parts)} / {len(note_parts)}')
    if len(set(BACKGROUND_HASHES))!=14:raise ValueError('A background image was reused')
    for log in LAYOUT_LOG:
        if log.get('box'):
            x,y,w,h=log['box']
            if min(x,y)<0 or x+w>W+.01 or y+h>H+.01:raise ValueError(f'Shape exceeds canvas: {log}')
    return {'slides':14,'nativeCharts':len(chart_parts),'speakerNotes':len(note_parts),'spokenSeconds':450,'uniqueBackgrounds':14,'editableText':True,'fontsEmbedded':False}


def render(pptx_path):
    executable='/workspace/.tools/soffice' if Path('/workspace/.tools/soffice').exists() else (shutil.which('libreoffice') or shutil.which('soffice'))
    if not executable:raise RuntimeError('LibreOffice headless is required to export the finished PDF/previews')
    profile=OUT/'lo-profile'
    command=[executable,'--headless',f'-env:UserInstallation={profile.as_uri()}','--convert-to','pdf','--outdir',str(OUT),str(pptx_path)]
    result=subprocess.run(command,check=True,capture_output=True,text=True,timeout=120)
    print(result.stdout.strip())
    preview_dir=OUT/'previews';preview_dir.mkdir(exist_ok=True)
    subprocess.run(['pdftoppm','-png','-r','96',str(pptx_path.with_suffix('.pdf')),str(preview_dir/'slide')],check=True,timeout=120)
    for number in range(1,15):
        unpadded=preview_dir/f'slide-{number}.png';target=preview_dir/f'slide-{number:02}.png'
        if unpadded.exists() and unpadded!=target:unpadded.replace(target)
    contact=Image.new('RGB',(4*480,4*270+80),(6,19,25))
    for index in range(14):
        thumb=Image.open(preview_dir/f'slide-{index+1:02}.png').convert('RGB').resize((480,270),Image.Resampling.LANCZOS)
        contact.paste(thumb,((index%4)*480,(index//4)*270))
    ImageDraw.Draw(contact).text((24,1096),'SHOCKTWIN · 14 SLIDES · 7:30 SCRIPT · 0:30 BUFFER',font=font_at('mono',14),fill='#ffffff')
    contact.save(preview_dir/'contact-sheet.png')
    return {'pdf':str(pptx_path.with_suffix('.pdf').relative_to(ROOT)),'previews':'presentations/previews/slide-01.png … slide-14.png','contactSheet':'presentations/previews/contact-sheet.png'}


prs=Presentation()
prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='ShockTwin — SAP HackFest 2026'
prs.core_properties.subject='Resilient Supply Chains · Detect, Simulate, Respond'
prs.core_properties.author='Francis Reuben R · Avaneesh · Affan · Vivek'
prs.core_properties.keywords='ShockTwin, CapacityMesh, SAP BTP, Illustrative demo'


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--render',action='store_true',help='Export matching PDF and 1920×1080 PNG previews through LibreOffice')
    args=parser.parse_args()
    input_check()
    ASSETS.mkdir(exist_ok=True)
    result=json.loads((ROOT/'docs/screenshots/simulation-result.json').read_text())
    assert len(result['baselineSeries'])==31 and result['daysToStockout']==18
    assert result['mitigatedDaysToStockout'] is None and result['assumptions']['switchDay']==3
    elapsed=0
    for item in DATA:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        background=prepare_background(item['number'])
        slide.shapes.add_picture(str(ASSETS/f"background-{item['number']:02}.png"),0,0,Inches(W),Inches(H))
        chrome(slide,background,item)
        eyebrow(slide,background,item['eyebrow'])
        headline(slide,background,item)
        fill_slide(slide,background,item,result)
        elapsed=notes(slide,item,elapsed)
    filename=OUT/'ShockTwin_SAP_HackFest.pptx'
    prs.save(filename)
    manifest=validate_file(filename)
    manifest.update({'pptx':str(filename.relative_to(ROOT)),'pptxSha256':sha256(filename),'canvasInches':[W,H],
                     'backgroundInputs':[{'slide':i,'path':f'presentations/backgrounds/slide-{i:02}.png','sha256':BACKGROUND_HASHES[i-1]} for i in range(1,15)],
                     'screenshots':str((ROOT/'docs/SCREENSHOT_PROVENANCE.json').relative_to(ROOT)),
                     'simulationInput':'docs/screenshots/simulation-result.json','composition':'Gaussian-blurred actual background crops with native text and borders',
                     'fontLicenseDirectory':'presentations/fonts','visualReview':'Required after every export; see QA.md'})
    if args.render:manifest.update(render(filename))
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    (OUT/'layout-report.json').write_text(json.dumps(LAYOUT_LOG,indent=2)+'\n')
    print(json.dumps({'pptx':str(filename),'slides':len(prs.slides),'spokenSeconds':elapsed,'nativeCharts':3,'rendered':args.render}))


if __name__=='__main__':main()
